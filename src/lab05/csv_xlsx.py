import csv
from pathlib import Path
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from typing import List, Any


def csv_to_xlsx(csv_path: str, xlsx_path: str) -> None:
    """Конвертирует CSV файл в XLSX формат"""
    cpath = Path(csv_path)
    if not cpath.exists():
        raise FileNotFoundError(f"Файл не найден: {csv_path}")

    # Чтение CSV файла
    rows: List[List[Any]] = []
    with cpath.open(encoding="utf-8") as f:
        csv_reader = csv.reader(f)
        for row in csv_reader:
            rows.append(row)

    if not rows:
        raise ValueError("CSV файл пуст")

    # Создание Excel файла
    wb = Workbook()
    ws = wb.active
    if ws is None:
        raise RuntimeError("Не удалось создать активный лист в Workbook")

    ws.title = "Data"

    # Запись данных
    for row in rows:
        ws.append(row)

    # Настройка автоширины колонок
    for col in ws.columns:
        if col:  # Проверка на пустую колонку
            max_length = 0
            for cell in col:
                try:
                    # Вычисляем длину текста в ячейке
                    cell_length = len(str(cell.value)) if cell.value else 0
                    max_length = max(max_length, cell_length)
                except:
                    continue

            # Устанавливаем ширину колонки (минимум 8 символов)
            col_idx = col[0].column
            if col_idx is not None:
                column_letter = get_column_letter(col_idx)
                ws.column_dimensions[column_letter].width = max(
                    8, min(max_length + 2, 50)
                )  # Ограничение максимум 50

    # Создание директории для выходного файла, если она не существует
    output_path = Path(xlsx_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    # Сохранение файла
    wb.save(xlsx_path)
    print(f"Файл успешно конвертирован: {xlsx_path}")


if __name__ == "__main__":
    try:
        csv_to_xlsx("data/samples/cities.csv", "data/out/cities.xlsx")
        csv_to_xlsx("data/samples/people.csv", "data/out/people.xlsx")
    except Exception as e:
        print(f"Ошибка при конвертации: {e}")
